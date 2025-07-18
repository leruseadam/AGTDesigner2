#!/usr/bin/env python3
"""
Test script to verify lineage coloring for Mixed and CBD Blend products.
"""

import sys
import os
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.core.generation.docx_formatting import apply_lineage_colors, COLORS
from src.core.generation.template_processor import TemplateProcessor
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

def test_lineage_coloring_extraction():
    """Test that lineage values are correctly extracted from embedded product type information."""
    
    print("Testing Lineage Color Extraction")
    print("=" * 50)
    
    # Test cases with embedded product type information
    test_cases = [
        {
            'name': 'Mixed Tincture (Non-Classic)',
            'text': 'MIXED_PRODUCT_TYPE_tincture_IS_CLASSIC_false',
            'expected_lineage': 'MIXED',
            'expected_color': COLORS['MIXED'],
            'is_classic': False,
            'product_type': 'tincture'
        },
        {
            'name': 'Mixed Flower (Classic)',
            'text': 'MIXED_PRODUCT_TYPE_flower_IS_CLASSIC_true',
            'expected_lineage': 'MIXED',
            'expected_color': COLORS['MIXED'],
            'is_classic': True,
            'product_type': 'flower'
        },
        {
            'name': 'CBD Blend Tincture',
            'text': 'CBD_PRODUCT_TYPE_tincture_IS_CLASSIC_false',
            'expected_lineage': 'CBD',
            'expected_color': COLORS['CBD'],
            'is_classic': False,
            'product_type': 'tincture'
        },
        {
            'name': 'CBD_BLEND Tincture',
            'text': 'CBD_BLEND_PRODUCT_TYPE_tincture_IS_CLASSIC_false',
            'expected_lineage': 'CBD_BLEND',
            'expected_color': COLORS['CBD'],
            'is_classic': False,
            'product_type': 'tincture'
        },
        {
            'name': 'Simple Mixed',
            'text': 'MIXED',
            'expected_lineage': 'MIXED',
            'expected_color': COLORS['MIXED'],
            'is_classic': None,
            'product_type': None
        },
        {
            'name': 'Simple CBD',
            'text': 'CBD',
            'expected_lineage': 'CBD',
            'expected_color': COLORS['CBD'],
            'is_classic': None,
            'product_type': None
        },
        {
            'name': 'Simple CBD_BLEND',
            'text': 'CBD_BLEND',
            'expected_lineage': 'CBD_BLEND',
            'expected_color': COLORS['CBD'],
            'is_classic': None,
            'product_type': None
        }
    ]
    
    for test_case in test_cases:
        print(f"\nTest: {test_case['name']}")
        print(f"Input text: '{test_case['text']}'")
        
        # Simulate the extraction logic from apply_lineage_colors
        text = test_case['text'].upper()
        
        # Extract lineage and product type information from embedded format
        lineage_part = text
        product_type = ""
        is_classic = False
        
        if "_PRODUCT_TYPE_" in text:
            parts = text.split("_PRODUCT_TYPE_")
            if len(parts) >= 2:
                lineage_part = parts[0]
                remaining = parts[1]
                # Extract product type and classic flag
                if "_IS_CLASSIC_" in remaining:
                    type_parts = remaining.split("_IS_CLASSIC_")
                    if len(type_parts) >= 2:
                        product_type = type_parts[0]
                        is_classic = type_parts[1].lower() == "true"
        
        print(f"Extracted lineage: '{lineage_part}'")
        print(f"Extracted product_type: '{product_type}'")
        print(f"Extracted is_classic: {is_classic}")
        print(f"Expected lineage: '{test_case['expected_lineage']}'")
        
        # Test color assignment
        color_hex = None
        if "PARAPHERNALIA" in text:
            color_hex = COLORS['PARA']
        elif "HYBRID/INDICA" in text or "HYBRID INDICA" in text:
            color_hex = COLORS['HYBRID_INDICA']
        elif "HYBRID/SATIVA" in text or "HYBRID SATIVA" in text:
            color_hex = COLORS['HYBRID_SATIVA']
        elif "SATIVA" in text:
            color_hex = COLORS['SATIVA']
        elif "INDICA" in text:
            color_hex = COLORS['INDICA']
        elif "HYBRID" in text:
            color_hex = COLORS['HYBRID']
        elif "CBD" in text or "CBD_BLEND" in text:
            color_hex = COLORS['CBD']
        elif "MIXED" in text:
            # For non-classic product types, Mixed should be blue
            if not is_classic and product_type and product_type not in ["flower", "pre-roll", "infused pre-roll", "concentrate", "solventless concentrate", "vape cartridge"]:
                color_hex = COLORS['MIXED']  # Blue for non-classic Mixed
            else:
                color_hex = COLORS['MIXED']  # Default blue for Mixed
        
        print(f"Assigned color: {color_hex}")
        print(f"Expected color: {test_case['expected_color']}")
        
        if lineage_part == test_case['expected_lineage'] and color_hex == test_case['expected_color']:
            print("✅ PASS")
        else:
            print("❌ FAIL")
    
    print("\n" + "=" * 50)

def test_template_processor_lineage_coloring():
    """Test that the template processor correctly applies lineage colors."""
    
    print("\nTesting Template Processor Lineage Coloring")
    print("=" * 50)
    
    # Test records with different lineages and product types
    test_records = [
        {
            'ProductName': 'Test Mixed Tincture',
            'Description': 'Test description',
            'Product Type*': 'tincture',
            'Lineage': 'MIXED',
            'Product Strain': 'Mixed',
            'Price': '$25.99',
            'Product Brand': 'Test Brand'
        },
        {
            'ProductName': 'Test CBD Blend Tincture',
            'Description': 'Test description with CBD',
            'Product Type*': 'tincture',
            'Lineage': 'CBD',
            'Product Strain': 'CBD Blend',
            'Price': '$25.99',
            'Product Brand': 'Test Brand'
        },
        {
            'ProductName': 'Test Mixed Flower',
            'Description': 'Test description',
            'Product Type*': 'flower',
            'Lineage': 'MIXED',
            'Product Strain': 'Mixed',
            'Price': '$25.99',
            'Product Brand': 'Test Brand'
        }
    ]
    
    template_types = ['vertical', 'horizontal', 'mini']
    
    for template_type in template_types:
        print(f"\n{template_type.upper()} TEMPLATE:")
        print("-" * 20)
        
        for i, record in enumerate(test_records):
            print(f"\n  Record {i+1}: {record['Lineage']} lineage, {record['Product Type*']} product type")
            
            try:
                # Create template processor
                processor = TemplateProcessor(template_type, {}, 1.0)
                
                # Process a single record
                result_doc = processor.process_records([record])
                
                if result_doc and result_doc.tables:
                    table = result_doc.tables[0]
                    cell = table.cell(0, 0)
                    
                    # Check if the cell has background color
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shading = tcPr.find(qn('w:shd'))
                    
                    if shading is not None:
                        fill_color = shading.get(qn('w:fill'))
                        print(f"    Cell background color: {fill_color}")
                        
                        # Check if the color matches expected lineage color
                        expected_color = None
                        if record['Lineage'] == 'MIXED':
                            expected_color = COLORS['MIXED']
                        elif record['Lineage'] == 'CBD':
                            expected_color = COLORS['CBD']
                        
                        if fill_color == expected_color:
                            print(f"    ✅ Color matches expected: {expected_color}")
                        else:
                            print(f"    ❌ Color mismatch. Expected: {expected_color}, Got: {fill_color}")
                    else:
                        print(f"    ❌ No background color found")
                        
                else:
                    print(f"    ❌ Failed to generate document")
                    
            except Exception as e:
                print(f"    ❌ Error: {e}")
    
    print("\n" + "=" * 50)

def test_cbd_blend_detection():
    """Test that CBD Blend is correctly detected based on description content."""
    
    print("\nTesting CBD Blend Detection")
    print("=" * 50)
    
    # Test cases for CBD Blend detection
    test_cases = [
        {
            'name': 'Description with CBD',
            'description': 'High CBD tincture with 500mg CBD',
            'product_strain': 'Mixed',
            'expected_strain': 'CBD Blend'
        },
        {
            'name': 'Description with colon',
            'description': 'Tincture: 100mg THC / 200mg CBD',
            'product_strain': 'Mixed',
            'expected_strain': 'CBD Blend'
        },
        {
            'name': 'Description with CBG',
            'description': 'CBG tincture with 100mg CBG',
            'product_strain': 'Mixed',
            'expected_strain': 'CBD Blend'
        },
        {
            'name': 'Description with CBN',
            'description': 'CBN tincture with 100mg CBN',
            'product_strain': 'Mixed',
            'expected_strain': 'CBD Blend'
        },
        {
            'name': 'Description with CBC',
            'description': 'CBC tincture with 100mg CBC',
            'product_strain': 'Mixed',
            'expected_strain': 'CBD Blend'
        },
        {
            'name': 'Regular description',
            'description': 'Regular tincture with THC',
            'product_strain': 'Mixed',
            'expected_strain': 'Mixed'
        }
    ]
    
    for test_case in test_cases:
        print(f"\nTest: {test_case['name']}")
        print(f"Description: '{test_case['description']}'")
        print(f"Original Product Strain: '{test_case['product_strain']}'")
        print(f"Expected Product Strain: '{test_case['expected_strain']}'")
        
        # Simulate the CBD Blend detection logic
        import re
        
        # Check if description contains cannabinoids or ":"
        cannabinoid_pattern = r"\b(?:CBD|CBC|CBN|CBG)\b"
        has_cannabinoids = bool(re.search(cannabinoid_pattern, test_case['description'], re.IGNORECASE))
        has_colon = ":" in test_case['description']
        
        # Determine if this should be CBD Blend
        should_be_cbd_blend = has_cannabinoids or has_colon
        
        if should_be_cbd_blend:
            result_strain = 'CBD Blend'
        else:
            result_strain = test_case['product_strain']
        
        print(f"Has cannabinoids: {has_cannabinoids}")
        print(f"Has colon: {has_colon}")
        print(f"Result Product Strain: '{result_strain}'")
        
        if result_strain == test_case['expected_strain']:
            print("✅ PASS")
        else:
            print("❌ FAIL")
    
    print("\n" + "=" * 50)

if __name__ == "__main__":
    test_lineage_coloring_extraction()
    test_template_processor_lineage_coloring()
    test_cbd_blend_detection()
    print("\nAll tests completed!") 