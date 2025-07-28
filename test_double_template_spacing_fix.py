#!/usr/bin/env python3
"""
Test script to debug spacing issue between brand and product strain in double template.
"""

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.core.generation.template_processor import TemplateProcessor
from docx import Document
import logging

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def test_double_template_spacing():
    """Test that brand and product strain have proper spacing in double template."""
    
    print("Testing Double Template Spacing")
    print("=" * 40)
    
    # Test data with brand and product strain
    test_records = [
        {
            'Product Name*': 'Test Product 1',
            'ProductBrand': 'Test Brand 1',
            'Product Brand': 'Test Brand 1',
            'Price': '$25.00',
            'Description': 'Test description 1',
            'Lineage': 'HYBRID',
            'THC_CBD': 'THC: 20% CBD: 2%',
            'WeightUnits': '3.5g',
            'ProductType': 'flower',
            'Product Strain': 'Test Strain 1'
        }
    ]
    
    try:
        # Test double template
        print("\n1. Testing Double Template Processing...")
        double_processor = TemplateProcessor('double', {}, 1.0)
        
        # Process records
        result = double_processor.process_records(test_records)
        assert result is not None, "Double template processing failed"
        
        # Check the first cell content
        table = result.tables[0]
        first_cell = table.cell(0, 0)
        cell_text = first_cell.text
        
        print(f"\n2. Cell content analysis:")
        print(f"   Raw cell text: '{cell_text}'")
        print(f"   Cell text length: {len(cell_text)}")
        
        # Check for line breaks
        lines = cell_text.split('\n')
        print(f"   Number of lines: {len(lines)}")
        for i, line in enumerate(lines):
            print(f"   Line {i}: '{line}'")
        
        # Check if brand and strain are properly separated
        if 'Test Brand 1' in cell_text and 'Test Strain 1' in cell_text:
            # Find the positions
            brand_pos = cell_text.find('Test Brand 1')
            strain_pos = cell_text.find('Test Strain 1')
            
            print(f"\n3. Content positions:")
            print(f"   Brand position: {brand_pos}")
            print(f"   Strain position: {strain_pos}")
            print(f"   Distance between: {strain_pos - brand_pos}")
            
            # Check if there's proper spacing between brand and strain
            text_between = cell_text[brand_pos + len('Test Brand 1'):strain_pos]
            print(f"   Text between brand and strain: '{text_between}'")
            
            # Check if there are line breaks between them
            has_line_break = '\n' in text_between
            print(f"   Has line break between: {has_line_break}")
            
            if not has_line_break:
                print("   ❌ ERROR: No line break between brand and strain")
                return False
            else:
                print("   ✓ Line break found between brand and strain")
        
        # Check paragraph structure
        print(f"\n4. Paragraph structure:")
        for i, para in enumerate(first_cell.paragraphs):
            print(f"   Paragraph {i}: '{para.text}'")
            print(f"   Paragraph runs: {len(para.runs)}")
            for j, run in enumerate(para.runs):
                print(f"     Run {j}: '{run.text}' (font size: {run.font.size})")
        
        print("\n" + "=" * 40)
        print("✓ SPACING TEST COMPLETED!")
        
        return True
        
    except Exception as e:
        print(f"\n❌ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    success = test_double_template_spacing()
    sys.exit(0 if success else 1) 