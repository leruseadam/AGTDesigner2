#!/usr/bin/env python3
"""
Test script to see how mini template works with multiple records.
"""

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.core.generation.template_processor import TemplateProcessor
from docx import Document

def test_mini_multiple_records():
    """Test mini template with multiple records."""
    
    print("Testing Mini Template with Multiple Records")
    print("=" * 50)
    
    # Create multiple test records
    test_records = []
    for i in range(5):  # Test with 5 records
        test_record = {
            'ProductName': f'Test Product {i+1}',
            'ProductBrand': f'Brand {i+1}',
            'Price': f'${(i+1)*10}.99',
            'Lineage': 'HYBRID',
            'Ratio_or_THC_CBD': 'THC: 25% CBD: 2%',
            'Description': f'Test Description {i+1}',
            'ProductStrain': '',
            'ProductType': 'Flower',
            'Vendor': f'Vendor {i+1}'
        }
        test_records.append(test_record)
    
    try:
        # Create processor
        processor = TemplateProcessor('mini', {}, 1.0)
        
        # Process records
        result_doc = processor.process_records(test_records)
        
        if result_doc and result_doc.tables:
            table = result_doc.tables[0]
            print(f"Result table: {len(table.rows)}x{len(table.columns)}")
            
            # Check all cells
            for row_idx in range(len(table.rows)):
                for col_idx in range(len(table.columns)):
                    cell = table.cell(row_idx, col_idx)
                    cell_text = cell.text.strip()
                    if cell_text:
                        print(f"Cell ({row_idx}, {col_idx}): '{cell_text}'")
                    else:
                        print(f"Cell ({row_idx}, {col_idx}): [empty]")
        else:
            print("❌ No result document or tables")
            
    except Exception as e:
        print(f"❌ ERROR: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_mini_multiple_records() 