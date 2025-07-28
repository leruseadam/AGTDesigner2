#!/usr/bin/env python3
"""
Test script to debug product strain font sizing in double template.
"""

import sys
import os
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from src.core.generation.template_processor import TemplateProcessor
from src.core.generation.unified_font_sizing import get_font_size_by_marker, get_font_size
from docx.shared import Pt
import logging

# Set up logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def test_strain_font_sizing():
    """Test that product strain gets 1pt font size in double template."""
    
    print("Testing Product Strain Font Sizing")
    print("=" * 40)
    
    # Test data
    test_strain = "Test Strain 1"
    
    try:
        print("\n1. Testing direct font sizing functions...")
        
        # Test get_font_size_by_marker for PRODUCTSTRAIN
        font_size = get_font_size_by_marker(test_strain, 'PRODUCTSTRAIN', 'double', 1.0)
        print(f"   get_font_size_by_marker('PRODUCTSTRAIN', 'double'): {font_size}")
        print(f"   Font size in points: {font_size.pt}")
        
        # Test get_font_size for strain field type
        font_size_direct = get_font_size(test_strain, 'strain', 'double', 1.0)
        print(f"   get_font_size('strain', 'double'): {font_size_direct}")
        print(f"   Font size in points: {font_size_direct.pt}")
        
        # Check if both return 1pt
        if font_size.pt == 1.0:
            print("   ✓ PRODUCTSTRAIN marker correctly returns 1pt font")
        else:
            print(f"   ❌ PRODUCTSTRAIN marker returns {font_size.pt}pt instead of 1pt")
            return False
        
        if font_size_direct.pt == 1.0:
            print("   ✓ Direct strain field type correctly returns 1pt font")
        else:
            print(f"   ❌ Direct strain field type returns {font_size_direct.pt}pt instead of 1pt")
            return False
        
        print("\n2. Testing template processor font sizing...")
        
        # Create template processor
        double_processor = TemplateProcessor('double', {}, 1.0)
        
        # Test the template-specific font sizing method
        template_font_size = double_processor._get_template_specific_font_size(test_strain, 'PRODUCTSTRAIN')
        print(f"   Template processor font size: {template_font_size}")
        print(f"   Font size in points: {template_font_size.pt}")
        
        if template_font_size.pt == 1.0:
            print("   ✓ Template processor correctly returns 1pt font")
        else:
            print(f"   ❌ Template processor returns {template_font_size.pt}pt instead of 1pt")
            return False
        
        print("\n3. Testing full template processing...")
        
        # Test with actual template processing
        test_records = [
            {
                'Product Name*': 'Test Product 1',
                'ProductBrand': 'Test Brand 1',
                'Product Brand': 'Test Brand 1',
                'Price': '$25.00',
                'Description': 'Test description 1',
                'Lineage': 'HYBRID',
                'THC_CBD': 'THC: 20% CBD: 2%',
                'WeightUnits': '3.5g',
                'ProductType': 'flower',
                'Product Strain': 'Test Strain 1'
            }
        ]
        
        result = double_processor.process_records(test_records)
        assert result is not None, "Template processing failed"
        
        # Check the first cell for font sizes
        table = result.tables[0]
        first_cell = table.cell(0, 0)
        
        print(f"\n4. Checking actual rendered font sizes...")
        
        # Check each paragraph and run for font sizes
        for i, para in enumerate(first_cell.paragraphs):
            if para.text.strip():
                print(f"   Paragraph {i}: '{para.text}'")
                for j, run in enumerate(para.runs):
                    font_size_pt = run.font.size.pt if run.font.size else "None"
                    print(f"     Run {j}: '{run.text}' - Font size: {font_size_pt}pt")
                    
                    # Check if this run contains strain content
                    if 'Test Strain 1' in run.text:
                        if font_size_pt == 1.0:
                            print(f"     ✓ Strain content has correct 1pt font size")
                        else:
                            print(f"     ❌ Strain content has {font_size_pt}pt font size instead of 1pt")
                            return False
        
        print("\n" + "=" * 40)
        print("✓ ALL FONT SIZING TESTS PASSED!")
        print("✓ Product strain correctly gets 1pt font size")
        
        return True
        
    except Exception as e:
        print(f"\n❌ TEST FAILED: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    success = test_strain_font_sizing()
    sys.exit(0 if success else 1) 